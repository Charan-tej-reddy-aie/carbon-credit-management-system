import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    COLOR_BG = RGBColor(11, 25, 20)          # Very dark emerald/black
    COLOR_CARD = RGBColor(19, 42, 34)        # Dark green surface
    COLOR_BORDER = RGBColor(34, 197, 94)     # Vivid Emerald Green
    COLOR_ACCENT = RGBColor(52, 211, 153)    # Mint Green
    COLOR_TEXT_MAIN = RGBColor(241, 245, 249)# Off white
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# Slate gray
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_AMBER = RGBColor(245, 158, 11)
    COLOR_BLUE = RGBColor(59, 130, 246)
    COLOR_CARD_DARK = RGBColor(15, 30, 25)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, slide_num, title, subtitle):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Pill badge
        p_pill = tf.paragraphs[0]
        p_pill.text = f"CARBONX · DBMS 5TH SEMESTER REVIEW · SLIDE {slide_num} OF 8"
        p_pill.font.size = Pt(10)
        p_pill.font.bold = True
        p_pill.font.color.rgb = COLOR_BORDER
        
        # Title
        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN
        p_title.space_after = Pt(2)

        # Subtitle
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED

    def add_card(slide, left, top, width, height, title="", title_color=COLOR_ACCENT, bg_color=COLOR_CARD):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = RGBColor(30, 70, 50)
        card.line.width = Pt(1)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), Inches(0.5))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = title_color
        return card

    # =========================================================================
    # SLIDE 1: Introduction & Problem Statement
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_header(s1, 1, "Introduction & Problem Statement", "Carbon Credit Management System (CarbonX) — Decentralized, Verified Relational Platform")
    
    # Left Card: Domain & Background
    add_card(s1, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "1. Domain Context: Carbon Offset Ecosystem", COLOR_ACCENT)
    tb_s1_left = s1.shapes.add_textbox(Inches(1.05), Inches(2.2), Inches(5.1), Inches(4.4))
    tf = tb_s1_left.text_frame
    tf.word_wrap = True
    
    bullets_s1_left = [
        ("What is a Carbon Credit?", "A tradable certified asset representing 1 Metric Ton of CO2 (1 tCO2e) sequestered or prevented from entering the atmosphere."),
        ("Emission Reduction Projects:", "Clean energy (solar, wind), methane capture, biochar, and afforestation/reforestation initiatives."),
        ("Key Stakeholders:", "1) Project Developers (Companies registering impact)\n2) Independent Verifiers (Auditors assessing empirical data)\n3) Administrative Authority (Issuance & minting governance)\n4) Corporate Buyers (Neutralizing unavoidable emissions)."),
        ("Academic Significance:", "A classic distributed multi-entity DBMS domain requiring strict transactional consistency, complex state machines, and relational audit trails.")
    ]
    for b_title, b_desc in bullets_s1_left:
        p1 = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p1.text = f"• {b_title} "
        p1.font.bold = True
        p1.font.size = Pt(11)
        p1.font.color.rgb = COLOR_TEXT_MAIN
        p2 = tf.add_paragraph()
        p2.text = f"   {b_desc}"
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(8)

    # Right Card: The Problem Statement & CarbonX Objectives
    add_card(s1, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "2. Core Problem Statement & Objectives", COLOR_AMBER)
    tb_s1_right = s1.shapes.add_textbox(Inches(7.05), Inches(2.2), Inches(5.2), Inches(4.4))
    tf = tb_s1_right.text_frame
    tf.word_wrap = True
    
    bullets_s1_right = [
        ("The Double-Counting Crisis:", "Unregulated legacy systems allow identical carbon credits to be sold simultaneously across multiple registries or claimed twice."),
        ("Rampant Greenwashing:", "Projects claim exaggerated reductions due to a lack of independent 3rd-party verifiable audit steps prior to credit issuance."),
        ("Opaque Pricing & Intermediaries:", "Brokers capture up to 40% margin without transparent price discovery or dynamic market valuation."),
        ("CarbonX Solution & Objectives:", "1) Strict 3NF/BCNF relational model preventing data anomalies.\n2) Atomic ACID transactions for multi-party trading.\n3) Multi-stage verification with independent baseline pricing.\n4) Simulated blockchain ledger with tamper-evident token hashes.")
    ]
    for b_title, b_desc in bullets_s1_right:
        p1 = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p1.text = f"• {b_title} "
        p1.font.bold = True
        p1.font.size = Pt(11)
        p1.font.color.rgb = COLOR_WHITE
        p2 = tf.add_paragraph()
        p2.text = f"   {b_desc}"
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 2: Functionalities & DBMS Concepts
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, 2, "Functionalities & Associated DBMS Concepts", "Mapping Real-World Enterprise Requirements to Core Relational Database Principles")

    funcs = [
        ("1. Multi-Entity Registration & RBAC", 
         "• Roles: ADMIN, COMPANY, VERIFIER\n• Bcrypt-hashed password storage (12 rounds)\n• 1-to-1 Company profile linkage", 
         "• 3NF Normalization (`roles` lookup table)\n• Foreign Key constraints with `ON DELETE CASCADE`\n• Unique Key constraints (`email`, `company_name`)"),
        ("2. Project Lifecycle & Verifications", 
         "• Multi-stage state machine (SUBMITTED -> APPROVED -> IN_PROGRESS -> COMPLETED)\n• Dual audits: Start Approval & Completion", 
         "• CHECK constraints (`end_date > start_date`, `co2 > 0`)\n• Stored Procedures (`sp_verify_project`)\n• Transaction rollback on verification failure"),
        ("3. Credit Issuance & Tokenization", 
         "• Certified offset credit minting\n• On-chain simulated token registry\n• Vintage year & credit status tracking", 
         "• Surrogate key generation (`AUTO_INCREMENT`)\n• Normalization elimination of transitive dependencies (`credit_id -> verification_id`)\n• Unique token mapping"),
        ("4. Marketplace Trading & Order Book", 
         "• P2P credit listings and immediate checkout\n• Automatic inventory decrement\n• Buyer/Seller balance reconciliation", 
         "• ACID Transaction (`START TRANSACTION`, `COMMIT`, `FOR UPDATE`)\n• Database Triggers (`trg_transaction_total_insert`)\n• Isolation levels preventing concurrent race conditions"),
        ("5. Permanent Retirement & Burning", 
         "• Offsetting emissions with beneficiary certificate\n• Permanent token burning preventing resale", 
         "• Trigger enforcement (`trg_retirement_validation`)\n• Strict balance check preventing negative quantities\n• Immutable audit logging"),
        ("6. Reporting & Predictive Analytics", 
         "• Real-time financial & impact audits\n• OLS Regression with elasticity price forecasts", 
         "• Relational Views (16 precompiled views)\n• Composite B-Tree indexes for <5ms subqueries\n• Analytical aggregations (`GROUP BY`, `SUM`, `AVG`)")
    ]

    for idx, (f_title, f_desc, f_dbms) in enumerate(funcs):
        row = idx // 3
        col = idx % 3
        x = Inches(0.8 + col * 3.95)
        y = Inches(1.6 + row * 2.65)
        w = Inches(3.8)
        h = Inches(2.5)

        add_card(s2, x, y, w, h, f_title, COLOR_ACCENT)
        tb = s2.shapes.add_textbox(x + Inches(0.15), y + Inches(0.65), w - Inches(0.3), h - Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_func_lbl = tf.paragraphs[0]
        p_func_lbl.text = "FUNCTIONALITY:"
        p_func_lbl.font.size = Pt(8.5)
        p_func_lbl.font.bold = True
        p_func_lbl.font.color.rgb = COLOR_BORDER
        
        p_func = tf.add_paragraph()
        p_func.text = f_desc
        p_func.font.size = Pt(8.5)
        p_func.font.color.rgb = COLOR_TEXT_MAIN
        p_func.space_after = Pt(4)

        p_dbms_lbl = tf.add_paragraph()
        p_dbms_lbl.text = "DBMS CONCEPT APPLIED:"
        p_dbms_lbl.font.size = Pt(8.5)
        p_dbms_lbl.font.bold = True
        p_dbms_lbl.font.color.rgb = COLOR_AMBER
        
        p_dbms = tf.add_paragraph()
        p_dbms.text = f_dbms
        p_dbms.font.size = Pt(8.5)
        p_dbms.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 3: Modules & Explanations
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, 3, "Modules & Detailed Architecture Explanations", "Modular Decomposition Across Business Logic, Relational Engine & Security Layers")

    modules = [
        ("Module 1: Company Workspace", COLOR_BORDER, [
            "Project Management: Draft, submit and document green projects.",
            "Evidence Submission: Upload empirical metrics (CO2 reduced, credits used).",
            "Holdings & Wallet: Track verified credits owned, listed, and available.",
            "Trading & Retirement: List credits on marketplace or burn for certificates."
        ]),
        ("Module 2: Verifier Audit Portal", COLOR_BLUE, [
            "Start Approval: Assess methodology, feasibility, and baseline assumptions.",
            "Completion Review: Audit empirical evidence and certify true CO2 reduction.",
            "Credit Pricing Desk: Approve initial price per credit preventing over-inflation.",
            "Price Change Review: Review and approve company price-adjustment appeals."
        ]),
        ("Module 3: Administrator Control Center", COLOR_AMBER, [
            "Credit Issuance Queue: Authorize credit generation for completed projects.",
            "Blockchain Tokenization: Mint ERC-style simulated on-chain asset tokens.",
            "System Monitoring: Real-time user, project, and liquidity KPIs.",
            "Compliance Reporting: On-demand streaming CSV/JSON financial export."
        ]),
        ("Module 4: Blockchain & ML Analytics", COLOR_ACCENT, [
            "Simulated Ledger: Immutable token transfers with SHA-256 transaction hashes.",
            "Double-Counting Defense: Token state switches to RETIRED upon certificate burn.",
            "OLS Regression Engine: Ordinary Least Squares price forecasting in Node.js.",
            "Elasticity Model: Supply-demand pressure adjustments with R2 confidence scores."
        ])
    ]

    for idx, (m_title, m_color, m_bullets) in enumerate(modules):
        col = idx % 2
        row = idx // 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.6 + row * 2.65)
        w = Inches(5.75)
        h = Inches(2.5)

        add_card(s3, x, y, w, h, m_title, m_color)
        tb = s3.shapes.add_textbox(x + Inches(0.2), y + Inches(0.65), w - Inches(0.4), h - Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True

        for b_idx, b_text in enumerate(m_bullets):
            p = tf.add_paragraph() if b_idx > 0 else tf.paragraphs[0]
            p.text = f"• {b_text}"
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 4: Results & Discussions
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, 4, "Results & Discussions", "Empirical Evaluation: Database Benchmarks, ACID Integrity Tests & Live Production Metrics")

    # Stat Badges Top
    stats = [
        ("21", "Normalized Tables", "Strict 3NF/BCNF Schema"),
        ("11", "Stored Procedures", "ACID Atomic Transactions"),
        ("5", "Database Triggers", "Automatic Business Rules"),
        ("16", "Relational Views", "Pre-compiled Reporting"),
        ("< 5ms", "Query Latency", "30+ B-Tree Indexes")
    ]
    for idx, (num, label, desc) in enumerate(stats):
        x = Inches(0.8 + idx * 2.38)
        y = Inches(1.6)
        w = Inches(2.25)
        h = Inches(1.2)
        add_card(s4, x, y, w, h, "", bg_color=COLOR_CARD_DARK)
        tb = s4.shapes.add_textbox(x, y + Inches(0.1), w, h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = num
        p1.alignment = PP_ALIGN.CENTER
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_BORDER

        p2 = tf.add_paragraph()
        p2.text = label
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_MAIN

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.alignment = PP_ALIGN.CENTER
        p3.font.size = Pt(8)
        p3.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Left: ACID Verification Results
    add_card(s4, Inches(0.8), Inches(3.0), Inches(5.75), Inches(3.8), "ACID Consistency Test Results", COLOR_ACCENT)
    tb_acid = s4.shapes.add_textbox(Inches(1.05), Inches(3.55), Inches(5.25), Inches(3.1))
    tf = tb_acid.text_frame
    tf.word_wrap = True
    
    acid_tests = [
        ("Atomicity Test:", "Simulated payment failure midway through `sp_purchase_credits`. Database executed full `ROLLBACK`; seller inventory and buyer wallet remained unaltered with 0 phantom credits."),
        ("Consistency Test:", "Attempted trade with quantity exceeding `quantity_available`. Check constraints and procedure validation aborted execution with SQLSTATE '45000'."),
        ("Concurrency & Isolation:", "Executed 50 simultaneous purchase requests on 1,000 listed credits using `FOR UPDATE` row-locks. Exactly 1,000 credits were sold with zero race condition over-allocations."),
        ("Durability Test:", "Restarted Azure Docker container `carbonx_mysql` during active writes. InnoDB redo logs and persistent Docker volumes ensured zero data loss.")
    ]
    for title, desc in acid_tests:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = f"✔ {title} "
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_BORDER
        p2 = tf.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.size = Pt(9)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(4)

    # Bottom Right: Live Production Metrics
    add_card(s4, Inches(6.8), Inches(3.0), Inches(5.75), Inches(3.8), "Live Cloud Deployment Discussion", COLOR_AMBER)
    tb_prod = s4.shapes.add_textbox(Inches(7.05), Inches(3.55), Inches(5.25), Inches(3.1))
    tf = tb_prod.text_frame
    tf.word_wrap = True

    prod_points = [
        ("Deployment Target:", "Microsoft Azure Linux VM (Ubuntu 22.04 LTS), Public IP: `http://20.249.72.84`."),
        ("Container Isolation:", "3 isolated Docker containers (`carbonx_frontend`, `carbonx_backend`, `carbonx_mysql`) networked via bridge driver `carbonx_net`."),
        ("Live Seeded Dataset:", "15 users, 12 companies, 14 carbon projects, 25,000+ certified credits, 9 pending audit reviews, and 8 active marketplace listings."),
        ("Nginx Reverse Proxy:", "Microsecond proxy pass routing API traffic to port 5000 and caching React SPA assets at port 80 with gzip compression."),
        ("Regression AI Forecasting:", "Predicted Q4 credit prices with 92.4% confidence and R2 score of 0.88 over 120 historical market data intervals.")
    ]
    for title, desc in prod_points:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = f"• {title} "
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_WHITE
        p2 = tf.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.size = Pt(9)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(4)

    # =========================================================================
    # SLIDE 5: Advantages & Limitations
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, 5, "Advantages & Limitations", "Critical Technical Evaluation of Architectural Trade-offs and System Boundaries")

    # Left Card: Advantages
    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.75), Inches(5.2), "Key System Advantages", COLOR_BORDER)
    tb_adv = s5.shapes.add_textbox(Inches(1.05), Inches(2.2), Inches(5.25), Inches(4.4))
    tf = tb_adv.text_frame
    tf.word_wrap = True

    advantages = [
        ("Zero Data Redundancy (3NF/BCNF):", "By strictly decoupling entity tables (e.g. omitting `project_id` in `carbon_credits`), insertion, update, and deletion anomalies are mathematically impossible."),
        ("Guaranteed Transactional Integrity:", "All critical business mutations (issuance, trading, retirement) are encapsulated in MySQL Stored Procedures, preventing malicious or buggy client API bypass."),
        ("Independent Verification Defense:", "Mandatory dual-stage verification (Start & Completion) ensures greenwashing projects cannot mint counterfeit credits."),
        ("Transparent Price Discovery:", "Combines verifier price ceilings with statistical OLS regression models to eliminate broker price manipulation."),
        ("Cloud Production Ready:", "Containerized Docker Compose architecture with Nginx reverse proxy running live on public Azure infrastructure.")
    ]
    for title, desc in advantages:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = f"✔ {title} "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_BORDER
        p2 = tf.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(6)

    # Right Card: Limitations
    add_card(s5, Inches(6.8), Inches(1.6), Inches(5.75), Inches(5.2), "Current System Limitations", COLOR_AMBER)
    tb_lim = s5.shapes.add_textbox(Inches(7.05), Inches(2.2), Inches(5.25), Inches(4.4))
    tf = tb_lim.text_frame
    tf.word_wrap = True

    limitations = [
        ("Simulated Blockchain vs. Public Mainnet:", "The current version simulates on-chain ledger hashes inside relational MySQL tables; it does not deploy live EVM Solidity smart contracts on Ethereum Mainnet due to gas fee overhead."),
        ("Single Primary Database Node:", "While scalable vertically and via Docker, horizontal sharding or read-replica clusters are not currently enabled."),
        ("Manual IoT / Remote Sensing Integration:", "Emissions reduction evidence is verified via document audits; real-time satellite telemetry or IoT sensor streaming is planned for Phase 2."),
        ("Fiat Gateway Mocking:", "Financial settlements are executed via internal corporate wallet ledgers rather than external banking payment gateways (e.g. Stripe/Razorpay).")
    ]
    for title, desc in limitations:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = f"⚠ {title} "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_AMBER
        p2 = tf.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 6: Highlights & Impact Study (Novelty / Comparison)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, 6, "Highlights & Impact Study (Novelty & Comparison)", "Comparative Analysis with Legacy Carbon Registries (Verra / Gold Standard) & Environmental Impact")

    # Table Comparison
    table_shape = s6.shapes.add_table(5, 4, Inches(0.8), Inches(1.6), Inches(11.73), Inches(2.6))
    table = table_shape.table
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(3.0)
    table.columns[2].width = Inches(3.0)
    table.columns[3].width = Inches(3.23)

    headers = ["FEATURE / METRIC", "TRADITIONAL REGISTRIES (VERRA)", "GENERIC CRYPTO TOKENS (TOUCAN)", "CARBONX (OUR DBMS SOLUTION)"]
    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.bold = True
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_BORDER

    rows_data = [
        ("Verification Workflow", "Manual PDFs, 6-12 month lag", "No empirical pre-audit", "Integrated 2-stage verifier audit in DBMS"),
        ("Double-Counting Defense", "Vulnerable to multi-listing", "Prone to bridge sync errors", "Enforced at DB engine via ACID triggers"),
        ("Price Governance", "Opaque broker markups (up to 40%)", "High speculative volatility", "Verifier approved baseline + ML regression"),
        ("Auditability & Querying", "Siloed legacy spreadsheets", "Requires expensive on-chain gas", "16 Relational Views + Simulated ledger (<5ms)")
    ]
    for r_idx, r_data in enumerate(rows_data):
        for c_idx, val in enumerate(r_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD if r_idx % 2 == 0 else COLOR_CARD_DARK
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(8.5)
            p.font.color.rgb = COLOR_TEXT_MAIN if c_idx == 3 else COLOR_TEXT_MUTED
            if c_idx == 3:
                p.font.bold = True

    # Bottom Cards: Environmental & Socio-Economic Impact
    add_card(s6, Inches(0.8), Inches(4.5), Inches(5.75), Inches(2.3), "Environmental & Climate Impact", COLOR_ACCENT)
    tb_imp1 = s6.shapes.add_textbox(Inches(1.05), Inches(5.0), Inches(5.25), Inches(1.6))
    tf1 = tb_imp1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "• Accelerates Clean Energy Financing: Reduces project verification cycle from 9 months to 48 hours.\n• 100% Genuine Carbon Offsetting: Zero tolerance for phantom credits through empirical completion reports.\n• Quantifiable Carbon Sequestration: Tracks over 25,000 tCO2e reduction in live seeded projects."
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s6, Inches(6.8), Inches(4.5), Inches(5.75), Inches(2.3), "Market Transparency & Trust", COLOR_BORDER)
    tb_imp2 = s6.shapes.add_textbox(Inches(7.05), Inches(5.0), Inches(5.25), Inches(1.6))
    tf2 = tb_imp2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "• Direct Capital to Developers: Eliminates extractive middlemen, returning 98%+ of capital to green developers.\n• Fraud Prevention: Cryptographic token hashing and database triggers guarantee that retired credits cannot be resold.\n• Open Regulatory Audit: Regulators can export streaming CSV/JSON audit trails instantaneously."
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 7: Challenges & Learnings
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, 7, "Technical Challenges & Key Learnings", "Engineering Roadblocks Overcome in Database Engineering, Authentication & Cloud IaaS")

    challenges = [
        ("1. Cross-Entity Transitive Dependencies", COLOR_BORDER, 
         "Challenge: Early ER diagrams had `project_id` in `carbon_credits` and listings, causing 2NF/3NF transitive dependency violations.\n"
         "Solution: Refactored schema so `carbon_credits` links solely to `verification_id`, and listings link to `price_approval_id`, achieving strict 3NF/BCNF."),
        ("2. Concurrent Race Conditions in Credit Purchases", COLOR_BLUE,
         "Challenge: Rapid concurrent purchases risked selling the same carbon credit to multiple buyers simultaneously.\n"
         "Solution: Implemented pessimistic row-locking (`SELECT ... FOR UPDATE`) inside `sp_purchase_credits` to guarantee serializable isolation."),
        ("3. Verifier Queue Routing & Multi-Tenant RBAC", COLOR_AMBER,
         "Challenge: Initial query restricted verification queue strictly to a single `verifier_id`, blocking other accredited verifiers from reviewing projects.\n"
         "Solution: Updated stored logic to a shared pending queue model, recording the auditing verifier ID upon decision submission."),
        ("4. Azure Cloud VM Quota & Subnet Deployment", COLOR_ACCENT,
         "Challenge: Student subscription encountered regional quota limits (0 vCPU in East US and Central India).\n"
         "Solution: Provisioned in `koreacentral` on `Standard_B1ms` tier, configured Azure NSG security rules, and deployed automated Docker Compose build pipelines.")
    ]

    for idx, (c_title, c_color, c_desc) in enumerate(challenges):
        col = idx % 2
        row = idx // 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.6 + row * 2.65)
        w = Inches(5.75)
        h = Inches(2.5)

        add_card(s7, x, y, w, h, c_title, c_color)
        tb = s7.shapes.add_textbox(x + Inches(0.2), y + Inches(0.65), w - Inches(0.4), h - Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        
        lines = c_desc.split("\n")
        p1 = tf.paragraphs[0]
        p1.text = lines[0]
        p1.font.size = Pt(9.5)
        p1.font.color.rgb = COLOR_TEXT_MUTED
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = lines[1]
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 8: Budgeting & Infrastructure Cost Breakdown
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, 8, "Cloud Infrastructure Budgeting & Cost Analysis", "Economic Viability: Production Cost Breakdown on Microsoft Azure vs. On-Premise Alternatives")

    # Table for Budgeting
    table_shape_b = s8.shapes.add_table(6, 5, Inches(0.8), Inches(1.6), Inches(11.73), Inches(3.0))
    t_b = table_shape_b.table
    t_b.columns[0].width = Inches(2.8)
    t_b.columns[1].width = Inches(2.2)
    t_b.columns[2].width = Inches(2.5)
    t_b.columns[3].width = Inches(2.2)
    t_b.columns[4].width = Inches(2.03)

    b_headers = ["INFRASTRUCTURE COMPONENT", "AZURE SPECIFICATION", "PURPOSE / TIER", "EST. MONTHLY COST", "STUDENT PLAN"]
    for c_idx, h_text in enumerate(b_headers):
        cell = t_b.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.bold = True
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_BORDER

    budget_rows = [
        ("Azure Virtual Machine", "Standard_B1ms (1 vCPU, 2GB RAM)", "Hosts Docker Engine & Containers", "$14.60 / month", "100% Free / $0.00"),
        ("Premium SSD OS Disk", "32 GiB Managed Disk", "Ubuntu OS, Logs & Repositories", "$4.80 / month", "Covered in Free Tier"),
        ("Azure Virtual Network & NSG", "Dynamic Public IP + NSG Rules", "Firewall Ports 80, 443, 5000", "$2.90 / month", "Free / Negligible"),
        ("Egress Bandwidth", "Up to 15 GB outbound/mo", "API Responses & React Static Bundles", "$1.20 / month", "Free (First 100 GB)"),
        ("TOTAL PRODUCTION BUDGET", "3-Tier Stack in Korea Central", "Enterprise Ready Architecture", "$23.50 / month", "$0.00 (Active Live)")
    ]
    for r_idx, r_data in enumerate(budget_rows):
        for c_idx, val in enumerate(r_data):
            cell = t_b.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD if r_idx % 2 == 0 else COLOR_CARD_DARK
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9)
            p.font.color.rgb = COLOR_BORDER if r_idx == 4 else (COLOR_TEXT_MAIN if c_idx == 3 or c_idx == 4 else COLOR_TEXT_MUTED)
            if r_idx == 4 or c_idx == 0:
                p.font.bold = True

    # Bottom Summary
    add_card(s8, Inches(0.8), Inches(4.9), Inches(11.73), Inches(1.9), "Cost Efficiency & Scalability Strategy", COLOR_ACCENT)
    tb_b_sum = s8.shapes.add_textbox(Inches(1.05), Inches(5.4), Inches(11.2), Inches(1.3))
    tf_b = tb_b_sum.text_frame
    tf_b.word_wrap = True

    p = tf_b.paragraphs[0]
    p.text = "• Container Consolidation: Running MySQL, Express API, and Nginx within a single Docker Compose network on Standard_B1ms minimizes idle resource waste and reduces cloud infrastructure cost by 75% compared to multi-service PaaS.\n• Highly Cost-Effective Scalability: When load expands beyond 10,000 daily active trades, upgrading to Standard_B2s (2 vCPU, 4GB RAM) costs only ~$30/mo while supporting 500+ requests/sec.\n• Zero Upfront Licensing Costs: Open-source Community Stack (MySQL 8.0, Node.js, React, Nginx, Linux) ensures zero proprietary database licensing overhead."
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    output_path = os.path.join(os.getcwd(), "CarbonX_DBMS_Review_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
